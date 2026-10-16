"""Import / export Excel."""
import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, date
from . import db


def _code(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s or None


def _num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _day(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str) and v.strip():
        for f in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), f).date().isoformat()
            except ValueError:
                pass
    return date.today().isoformat()


def _table(ws):
    rs = list(ws.iter_rows(values_only=True))
    if not rs:
        return []
    head = [str(h).strip() if h is not None else "" for h in rs[0]]
    out = []
    for r in rs[1:]:
        if any(c is not None for c in r):
            out.append({head[i]: r[i] for i in range(min(len(head), len(r))) if head[i]})
    return out


def import_workbook(path, replace=True):
    import warnings
    warnings.simplefilter("ignore")
    wb = openpyxl.load_workbook(path, data_only=True)
    report = {"books": 0, "packs": 0, "lots": 0, "sales": 0, "reservations": 0, "skipped": []}
    names = {n.lower(): n for n in wb.sheetnames}

    def sheet(*keys):
        for n in wb.sheetnames:
            if all(k in n.lower() for k in keys):
                return wb[n]
        return None

    if os.path.exists(db.DB_PATH):
        shutil.copy(db.DB_PATH, os.path.join(db.DATA_DIR, f"avant_import_{datetime.now():%Y%m%d_%H%M%S}.db"))

    with db.tx() as c:
        if replace:
            for t in ("allocations", "sale_lines", "sales", "adjustments", "lots", "pack_items",
                      "settlements", "reservations", "books"):
                c.execute(f"DELETE FROM {t}")
        by_code, by_unique = {}, {}
        std = sheet("catalogue", "std")
        n_auto = 0
        if std:
            for r in _table(std):
                title = r.get("Titre")
                if not title:
                    continue
                code = _code(r.get("Code"))
                if not code:
                    n_auto += 1
                    code = f"SANS-CODE-{n_auto}"
                if code in by_code:
                    report["skipped"].append(f"Doublon de code ignoré : {code}")
                    continue
                vol = r.get("Volume")
                vol = _code(vol) if vol is not None else None
                unique = r.get("Unique") or " - ".join(
                    [x for x in [r.get("Série"), f"Vol. {vol}" if vol else None, title, r.get("Type")] if x])
                cur = c.execute("""INSERT INTO books(code,title,series,volume,author,publisher,category,type,location,price)
                    VALUES(?,?,?,?,?,?,?,?,?,?)""",
                    (code, title, r.get("Série"), vol, r.get("Auteur"), r.get("Editeur"),
                     r.get("Catégorie") or "Standard", r.get("Type"), r.get("Emplacement"), _num(r.get("Prix"))))
                by_code[code] = cur.lastrowid
                by_unique[str(unique).strip()] = cur.lastrowid
                report["books"] += 1
        pk = sheet("catalogue", "pack")
        if pk:
            for r in _table(pk):
                title, code = r.get("Titre"), _code(r.get("Code"))
                if not title or not code or code in by_code:
                    continue
                cur = c.execute("INSERT INTO books(code,title,category,price,is_pack) VALUES(?,?,?,?,1)",
                                (code, str(title).strip(), "Pack", _num(r.get("Prix"))))
                pid = cur.lastrowid
                by_code[code] = pid
                for k in ("Titre 1", "Titre 2", "Titre 3", "Titre 4", "Titre 5"):
                    t = r.get(k)
                    bid = by_unique.get(str(t).strip()) if t else None
                    if t and bid:
                        c.execute("INSERT INTO pack_items(pack_id,book_id,qty) VALUES(?,?,1) "
                                  "ON CONFLICT(pack_id,book_id) DO UPDATE SET qty=qty+1", (pid, bid))
                    elif t:
                        report["skipped"].append(f"Pack {code} : composant introuvable « {t} »")
                report["packs"] += 1
        owners = {o["name"]: o["id"] for o in db.rows(c, "SELECT id,name FROM owners")}
        liv = sheet("livraison")
        if liv:
            for r in _table(liv):
                code = _code(r.get("Code"))
                bid = by_code.get(code) or by_unique.get(str(r.get("Titre") or "").strip())
                q = int(_num(r.get("Quantité Reçu")))
                if not bid or q == 0:
                    if code or r.get("Titre"):
                        report["skipped"].append(f"Livraison ignorée : {code or r.get('Titre')}")
                    continue
                on = (r.get("Propriétaire") or "Cité Royale").strip()
                if on not in owners:
                    owners[on] = c.execute("INSERT INTO owners(name,kind,reverse_pct,priority) VALUES(?,'locale',0,9)",
                                           (on,)).lastrowid
                c.execute("INSERT INTO lots(book_id,owner_id,date,qty_in,note) VALUES(?,?,?,?,?)",
                          (bid, owners[on], _day(r.get("Date de Livraison")), q, r.get("Observation")))
                report["lots"] += 1
        ven = sheet("vente")
        if ven:
            groups = {}
            for r in _table(ven):
                code = _code(r.get("Code"))
                bid = by_code.get(code) or by_unique.get(str(r.get("Titre") or "").strip())
                q = int(_num(r.get("Quantité")))
                if not bid or q <= 0:
                    if code or r.get("Titre"):
                        report["skipped"].append(f"Vente ignorée : {code or r.get('Titre')}")
                    continue
                price = _num(r.get("Prix Unitaire"), None)
                groups.setdefault((_day(r.get("Date")), r.get("Moyen") or "Espèces"), []).append(
                    {"book_id": bid, "qty": q, "unit_price": price})
            for (d, m), lines in sorted(groups.items()):
                db.create_sale(c, d, m, "Import Excel", lines, manual=False)
                report["sales"] += 1
        cmd = sheet("commande")
        if cmd:
            for r in cmd.iter_rows(values_only=True):
                if r[1] and r[2]:
                    c.execute("INSERT INTO reservations(name,title,phone,date) VALUES(?,?,?,?)",
                              (r[1], r[2], _code(r[4]) if len(r) > 4 else None, date.today().isoformat()))
                    report["reservations"] += 1
    return report


def _sheet(wb, title, headers, data, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in data:
        ws.append(row)
    fill = PatternFill("solid", fgColor="2B3A67")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for i, h in enumerate(headers, 1):
        w = max([len(str(h))] + [len(str(r[i - 1])) for r in data[:200] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(60, w + 3)
    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = ws.dimensions
    return ws


def export_workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with db.tx() as c:
        books = db.rows(c, "SELECT * FROM books WHERE is_pack=0 ORDER BY series,title,volume")
        _sheet(wb, "Catalogue", ["Code", "Série", "Titre", "Volume", "Auteur", "Éditeur", "Catégorie", "Type",
                                 "Emplacement", "Prix", "Stock actuel", "Seuil"],
               [[b["code"], b["series"], b["title"], b["volume"], b["author"], b["publisher"], b["category"],
                 b["type"], b["location"], b["price"], db.book_stock(c, b["id"]), b["min_stock"]] for b in books])
        packs = db.rows(c, "SELECT * FROM books WHERE is_pack=1")
        data = []
        for p in packs:
            comp = db.rows(c, "SELECT bk.title,pi.qty FROM pack_items pi JOIN books bk ON bk.id=pi.book_id WHERE pi.pack_id=?", (p["id"],))
            data.append([p["code"], p["title"], p["price"], db.pack_stock(c, p["id"]),
                         " + ".join(f"{x['qty']}× {x['title']}" for x in comp)])
        _sheet(wb, "Packs", ["Code", "Titre", "Prix", "Stock possible", "Composition"], data)
        lots = db.rows(c, f"""SELECT l.date,b.code,b.title,o.name owner,l.qty_in,l.unit_cost,
            COALESCE(l.reverse_pct,o.reverse_pct) pct,({db.LOT_REMAINING}) remaining,l.note
            FROM lots l JOIN books b ON b.id=l.book_id JOIN owners o ON o.id=l.owner_id ORDER BY l.date""")
        _sheet(wb, "Livraisons", ["Date", "Code", "Titre", "Propriétaire", "Quantité reçue", "Coût unitaire",
                                  "% reversé", "Restant", "Observation"],
               [list(x.values()) for x in lots])
        sales = db.rows(c, """SELECT s.date,s.payment,b.code,b.title,l.qty,l.unit_price,l.qty*l.unit_price total
            FROM sale_lines l JOIN sales s ON s.id=l.sale_id JOIN books b ON b.id=l.book_id ORDER BY s.date,s.id""")
        _sheet(wb, "Ventes", ["Date", "Moyen", "Code", "Titre", "Quantité", "Prix unitaire", "Prix total"],
               [list(x.values()) for x in sales])
        rev = db.rows(c, """SELECT o.name,COALESCE(SUM(a.revenue),0) rev,COALESCE(SUM(a.due),0) due
            FROM owners o LEFT JOIN allocations a ON a.owner_id=o.id GROUP BY o.id""")
        data = []
        for r in rev:
            paid = c.execute("SELECT COALESCE(SUM(amount),0) s FROM settlements WHERE owner_id=(SELECT id FROM owners WHERE name=?)", (r["name"],)).fetchone()["s"]
            data.append([r["name"], r["rev"], r["due"], paid, r["due"] - paid, r["rev"] - r["due"]])
        _sheet(wb, "Reversements", ["Propriétaire", "CA généré", "Dû", "Déjà réglé", "Reste à régler", "Marge gardée"], data)
        _sheet(wb, "Réservations", ["Nom", "Téléphone", "Titre", "Qté", "Date", "Statut", "Note"],
               [[r["name"], r["phone"], r["title"], r["qty"], r["date"], r["status"], r["note"]]
                for r in db.rows(c, "SELECT * FROM reservations")])
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)
    return path
